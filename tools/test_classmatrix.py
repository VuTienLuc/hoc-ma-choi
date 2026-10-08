"""Kiểm tra bảng tổng hợp tiến độ cả lớp + xuất Excel."""
import asyncio, json, pathlib, sys, tempfile
from playwright.async_api import async_playwright
ROOT = pathlib.Path(__file__).resolve().parent.parent

async def prepare(page, teacher):
    acc = {"token": "t", "name": "Thầy Lực" if teacher else "An", "lop": "GV" if teacher else "10A1", "user": "gv" if teacher else "an"}
    await page.add_init_script(f"localStorage.setItem('hoctap:session', JSON.stringify({json.dumps(acc)}))")
    await page.goto((ROOT / "index.html").as_uri())
    await page.wait_for_timeout(700)
    await page.evaluate("() => { location.hash = '#/lop10'; }")
    await page.wait_for_timeout(500)

async def main():
    bad = []
    async with async_playwright() as pw:
        br = await pw.chromium.launch()
        ctx = await br.new_context(accept_downloads=True, viewport={"width": 1024, "height": 768})
        page = await ctx.new_page(); errs = []
        page.on("pageerror", lambda e: errs.append(str(e)))
        await prepare(page, True)
        if await page.locator("[data-class-matrix]").count() != 1:
            bad.append("GV không thấy nút Thống kê lớp")
        for fallback in (False, True):
            await page.evaluate("""(fb) => {
              const L = App.grades.find(g=>g.id==='lop10').lessons, a=L[0].id, b=L[1].id;
              window.__calls=[];
              Account.call = async (act, body) => { window.__calls.push(act);
                if(act==='gradeProgress' && !fb) return {ok:true, grade:'lop10', classes:[
                  {lop:'10A1', rows:[{name:'An',user:'an',joined:true,last:'',p:{}},
                    {name:'Bình',user:'binh',joined:true,last:'08/10/2026',p:{[a]:[3,2,null]}},
                    {name:'Chi',user:'chi',joined:true,last:'08/10/2026',p:{[a]:[3,3,2],[b]:[1,null,null]}}]},
                  {lop:'10A2', rows:[{name:'Dung',user:'dung',joined:false,last:'',p:{}}]}]};
                if(act==='gradeProgress') return {ok:false,msg:'Hành động không hợp lệ'};
                const done = body.lesson===a;
                return {ok:true, classes:[{lop:'10A1', rows:[
                  {name:'An',user:'an',joined:true,levels:[null,null,null],completed:0,last:''},
                  {name:'Chi',user:'chi',joined:true,levels:done?[3,3,2]:[1,null,null],completed:done?3:1,last:'08/10/2026'}]}]};
              };
            }""", fallback)
            await page.locator("[data-class-matrix]").click()
            await page.wait_for_selector(".cm-tb", timeout=8000)
            tag = "dự phòng" if fallback else "gradeProgress"
            names = await page.locator(".cm-tb tbody .cm-name").all_inner_texts()
            if not any("Chi" in n for n in names) or not any("An" in n for n in names):
                bad.append(f"[{tag}] bảng thiếu học sinh: {names}")
            if await page.locator(".cm-c.done").count() < 1: bad.append(f"[{tag}] không có ô xanh (đủ 3 mức)")
            if await page.locator(".cm-c.todo").count() < 1: bad.append(f"[{tag}] không có ô đỏ (chưa làm)")
            await page.locator(".cm-les").first.click()
            det = await page.locator(".cm-detail").inner_text()
            if "An" not in det or "Chưa" not in det: bad.append(f"[{tag}] khung chi tiết bài không liệt kê em chưa làm: {det[:80]}")
            if not fallback:
                if "lessonStatus" in await page.evaluate("window.__calls"): bad.append("Gọi lessonStatus dù gradeProgress chạy được")
                if await page.locator("#cmClass option").count() < 2: bad.append("Thiếu chọn lớp 10A2")
            async with page.expect_download() as dl:
                await page.locator("#cmXlsx").click()
            path = pathlib.Path(tempfile.mkdtemp()) / "t.xlsx"
            await (await dl.value).save_as(path)
            try:
                from openpyxl import load_workbook
                wb = load_workbook(path); ws = wb.worksheets
                if len(ws) != 2: bad.append(f"[{tag}] xlsx có {len(ws)} sheet, cần 2")
                else:
                    flat = [str(c.value) for r in ws[0].iter_rows() for c in r if c.value is not None]
                    flat2 = [str(c.value) for r in ws[1].iter_rows() for c in r if c.value is not None]
                    if "Chi" not in " ".join(flat): bad.append(f"[{tag}] sheet 1 thiếu tên học sinh")
                    if "An" not in " ".join(flat2): bad.append(f"[{tag}] sheet 2 thiếu danh sách chưa làm")
                    fills = {c.fill.fgColor.rgb for r in ws[0].iter_rows() for c in r if c.fill and c.fill.fill_type}
                    if len(fills) < 2: bad.append(f"[{tag}] xlsx không tô màu trạng thái")
            except Exception as e:
                bad.append(f"[{tag}] không mở được xlsx: {e}")
            if not fallback:
                ov = await page.evaluate("(()=>{const c=document.querySelector('.cm-card').getBoundingClientRect();return [c.left,c.right,innerWidth,document.documentElement.scrollWidth]})()")
                if ov[0] < -1 or ov[1] > ov[2] + 1 or ov[3] > ov[2] + 1: bad.append(f"Bảng tràn ngang: {ov}")
            await page.locator("[data-cm-close]").click()
            if await page.locator(".cm-overlay").count(): bad.append("Không đóng được bảng")
        bad += ["Lỗi trang: " + e for e in errs]
        s = await br.new_page(); await prepare(s, False)
        if await s.locator("[data-class-matrix]").count(): bad.append("Học sinh vẫn thấy nút Thống kê lớp")
        await br.close()
    print("Bảng tổng hợp cả lớp + Excel: " + ("ĐẠT" if not bad else "KHÔNG ĐẠT"))
    for b in bad: print(" ✗", b)
    sys.exit(1 if bad else 0)
asyncio.run(main())
